"""Build the editable Word handbook from docs/REPOSITORY_GUIDE.md.

Requires python-docx. Run from any directory; paths are resolved from this file.
This documentation utility does not import or run the engineering models.
Open the result in Word and update its contents table after regenerating it.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/REPOSITORY_GUIDE.md"
OUTPUT = ROOT / "docs/BEAST_Repository_Guide.docx"


def element(tag, **attributes):
    """Create a Word XML element for properties python-docx does not expose."""
    value = OxmlElement(tag)
    for key, item in attributes.items():
        value.set(qn("w:" + key), str(item))
    return value


def field(paragraph, instruction):
    """Insert an editable Word field, such as a page number or contents table."""
    paragraph._p.append(element("w:r"))
    run = paragraph._p[-1]
    run.append(element("w:fldChar", fldCharType="begin"))
    code = element("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    run.append(code)
    run.append(element("w:fldChar", fldCharType="end"))


def inline(paragraph, text):
    """Translate just bold and inline code; leave ordinary prose as Word text."""
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(part)


def configure(document):
    """Use a restrained Letter-size handbook layout with readable body text."""
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    section.header_distance = section.footer_distance = Inches(0.3)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        # Some Word templates carry a title rule; the handbook uses plain titles.
        for border in list(style.element.xpath(".//w:pBdr")):
            border.getparent().remove(border)
    body = document.styles["Normal"]
    body.font.size = Pt(11)
    body.paragraph_format.line_spacing = 1.14
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.widow_control = True
    document.styles["Title"].font.size = Pt(26)
    document.styles["Title"].font.bold = True
    document.styles["Subtitle"].font.size = Pt(13)
    for name, size in [("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11.5)]:
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(16)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = "BEAST Repository Reading Guide"
    header.runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Page ")
    field(footer, "PAGE")
    footer.add_run(" of ")
    field(footer, "NUMPAGES")
    for run in footer.runs:
        run.font.size = Pt(9)
    document.settings._element.append(element("w:updateFields", val="true"))


def cells(line):
    # A literal vertical bar inside a Markdown table is escaped as \|.
    return [cell.strip().replace("\\|", "|") for cell in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def add_table(document, lines):
    """Keep repeated records in a real Word table with a repeating header."""
    rows = [cells(line) for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r"[:\- ]+", c) for c in row)]
    columns = len(rows[0])
    heading = rows[0][0]
    if columns == 2:
        widths = [4.2, 2.8] if heading == "Command" else [3.25, 3.75]
    elif heading == "TypeScript expression":
        widths = [2.05, 2.05, 2.9]
    elif heading == "Record":
        widths = [1.6, 2.15, 3.25]
    else:
        widths = [2.1, 2.45, 2.45]
    table = document.add_table(rows=0, cols=columns)
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    properties = table._tbl.tblPr
    borders = element("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        borders.append(element("w:" + edge, val="single", sz="4", color="D9D9D9"))
    properties.append(borders)
    for index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(element("w:cantSplit"))
        if index == 0:
            row._tr.get_or_add_trPr().append(element("w:tblHeader"))
        for cell, text, width in zip(row.cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = element("w:tcMar")
            for edge in ["top", "bottom", "left", "right"]:
                margins.append(element("w:" + edge, w="90", type="dxa"))
            props.append(margins)
            props.append(element("w:shd", fill="E7EEF5" if index == 0 else "FFFFFF"))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.08
            paragraph.paragraph_format.keep_with_next = index == 0
            inline(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(9.5 if run.font.name == "Consolas" else 10)
                if index == 0:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    document = Document()
    configure(document)
    document.core_properties.title = "BEAST Repository Reading Guide"
    document.core_properties.subject = "Engineering code navigation and maintenance"
    document.core_properties.author = "BEAST"
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    title_paragraph = document.add_paragraph(lines[0][2:], "Title")
    for border in list(title_paragraph._p.xpath(".//w:pBdr")):
        border.getparent().remove(border)
    document.add_paragraph(lines[2], "Subtitle")
    document.add_paragraph(lines[4])
    document.add_paragraph("Read the calculation modules as connected engineering worksheets. This handbook explains their inputs, units, outputs and links to the website, using Python comparisons where the TypeScript syntax is unfamiliar.")
    document.add_paragraph("The annotated code and this handbook are intended to help you trace results and make focused, verifiable changes. Start with the physical models, then read the interface and server connections.")
    document.add_page_break()
    title = document.add_paragraph("Contents")
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(16)
    field(document.add_paragraph(), 'TOC \\o "1-1" \\h \\z \\u')
    document.add_page_break()
    index = 6
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("|"):
            block = []
            while index < len(lines) and lines[index].startswith("|"):
                block.append(lines[index]); index += 1
            add_table(document, block)
            continue
        if line.startswith("```"):
            code = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index]); index += 1
            for position, text in enumerate(code):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.17)
                paragraph.paragraph_format.space_after = Pt(1 if position < len(code) - 1 else 8)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.paragraph_format.keep_with_next = position < len(code) - 1
                run = paragraph.add_run(text)
                run.font.name, run.font.size = "Consolas", Pt(9)
            index += 1
            continue
        if line.startswith("## "):
            document.add_paragraph(line[3:], "Heading 1")
        elif line.startswith("### "):
            document.add_paragraph(line[4:], "Heading 2")
        elif re.match(r"^\d+\. ", line):
            # Preserve the manuscript's list numbers; Word must not continue a
            # previous chapter's independent procedure after opening/saving.
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.2)
            paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(0.25))
            number, text = line.split(". ", 1)
            paragraph.add_run(number + ".\t")
            inline(paragraph, text)
        else:
            paragraph = document.add_paragraph()
            inline(paragraph, line)
            following = next((value for value in lines[index + 1:] if value.strip()), "")
            if following.startswith(("```", "|")):
                paragraph.paragraph_format.keep_with_next = True
        index += 1
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
